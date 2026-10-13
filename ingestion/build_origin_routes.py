"""Import a route workbook into the dashboard's non-PII route catalog."""

from __future__ import annotations

import json
import sys
from pathlib import Path

import openpyxl


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "dist" / "warehouse-carriers.json"
HEADERS = ["始发仓库", "物流公司", "收货城市", "收货省份", "始发区域"]


def load_routes(source: Path) -> list[dict[str, str]]:
    workbook = openpyxl.load_workbook(source, read_only=True, data_only=True)
    sheet = workbook.active
    headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    if headers[:len(HEADERS)] != HEADERS:
        raise ValueError(f"表头不符合预期：{headers}")
    routes = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        values = [str(value).strip() if value is not None else "" for value in row[:5]]
        if not any(values):
            continue
        if not all(values):
            raise ValueError(f"发现字段不完整的线路记录：{values}")
        routes.append(dict(zip(("warehouse", "carrier", "city", "province", "region"), values)))
    if not routes:
        raise ValueError("表格中没有线路记录")
    return routes


def carrier_map(routes: list[dict[str, str]]) -> dict[str, list[str]]:
    mapping: dict[str, set[str]] = {}
    for route in routes:
        mapping.setdefault(route["warehouse"], set()).add(route["carrier"])
    return {warehouse: sorted(carriers) for warehouse, carriers in sorted(mapping.items())}


def main() -> None:
    if len(sys.argv) != 2:
        raise SystemExit("用法：python ingestion/build_origin_routes.py <线路信息.xlsx>")
    routes = load_routes(Path(sys.argv[1]).expanduser())
    mapping = carrier_map(routes)
    existing = json.loads(OUTPUT.read_text(encoding="utf-8")) if OUTPUT.exists() else {}
    catalog = {"source": "origin-route-and-warehouse-region-workbooks",
               "warehouse_carriers": mapping}
    if existing.get("warehouse_regions"):
        catalog["warehouse_regions"] = existing["warehouse_regions"]
    OUTPUT.write_text(json.dumps(catalog, ensure_ascii=False, separators=(",", ":")) + "\n",
                      encoding="utf-8")
    print(f"已生成 {OUTPUT}：{len(mapping)} 个始发仓的承运商关系")


if __name__ == "__main__":
    main()
