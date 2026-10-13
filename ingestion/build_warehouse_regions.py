"""Merge a warehouse-to-region workbook into the dashboard filter catalog."""

from __future__ import annotations

import json
import sys
from pathlib import Path

import openpyxl


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "dist" / "warehouse-carriers.json"


def load_warehouse_regions(source: Path) -> dict[str, str]:
    workbook = openpyxl.load_workbook(source, read_only=True, data_only=True)
    sheet = workbook.active
    header = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
    positions = {str(value).strip(): index for index, value in enumerate(header)
                 if value is not None}
    required = ("仓库名称", "仓区域")
    if any(name not in positions for name in required):
        raise ValueError(f"表头不符合预期，需要字段：{'、'.join(required)}")

    regions: dict[str, str] = {}
    for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        warehouse = str(row[positions["仓库名称"]] or "").strip()
        region = str(row[positions["仓区域"]] or "").strip()
        if not warehouse and not region:
            continue
        if not warehouse or not region:
            raise ValueError(f"第 {row_number} 行仓库或区域为空")
        prior = regions.get(warehouse)
        if prior and prior != region:
            raise ValueError(f"仓库「{warehouse}」重复映射到「{prior}」和「{region}」")
        regions[warehouse] = region
    if not regions:
        raise ValueError("工作簿中没有仓库区域映射")
    return regions


def write_catalog(regions: dict[str, str], output: Path = OUTPUT) -> None:
    catalog = json.loads(output.read_text(encoding="utf-8")) if output.exists() else {}
    catalog["source"] = "origin-route-and-warehouse-region-workbooks"
    catalog["warehouse_regions"] = regions
    output.write_text(json.dumps(catalog, ensure_ascii=False, separators=(",", ":")) + "\n",
                      encoding="utf-8")


def main() -> None:
    if len(sys.argv) != 2:
        raise SystemExit("用法：python ingestion/build_warehouse_regions.py <仓-区域映射.xlsx>")
    regions = load_warehouse_regions(Path(sys.argv[1]).expanduser())
    write_catalog(regions)
    print(f"已更新区域映射：{len(regions)} 个仓，{len(set(regions.values()))} 个统仓区域")


if __name__ == "__main__":
    main()
