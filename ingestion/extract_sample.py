#!/usr/bin/env python3
"""Extract the SOO fact rows from the reference workbook into dashboard JSON.

This is deliberately a small, deterministic MVP ingestion step. It keeps the
source workbook untouched, filters formula-only blank rows, normalizes dates
and numbers, and records quality flags instead of silently turning bad values
into zeros.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
from datetime import date, datetime, time, timedelta
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo

import openpyxl

BUSINESS_TZ = ZoneInfo("Asia/Shanghai")
SOURCE_COLUMNS = 55  # A:BC; BD:BH are workbook formulas, not TMS fields.
REQUIRED_COLUMNS = {"域", "SOO订单号", "生产日期", "装载完成时间", "实际提货日期", "状态", "soo下发件数"}
PICKUP_PENDING_STATUSES = frozenset({"开放", "委托", "确认"})


CHANNEL_RULES = [
    ("京东", "京东"),
    ("天猫", "天猫"),
    ("线下经销", "线下经销"),
    ("子行", "子行科技"),
]


def clean_text(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def is_picked_up_status(value: Any) -> bool:
    """TMS logistics status is the source of truth for pickup completion.

    A missing status is not evidence of completion; all other non-empty source
    statuses count as picked up per the agreed business rule.
    """
    status = clean_text(value)
    return status is not None and status != "未知" and status not in PICKUP_PENDING_STATUSES


def parse_number(value: Any) -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        number = float(value)
        return number if math.isfinite(number) else None
    text = str(value).strip().replace(",", "")
    if not text:
        return None
    try:
        number = float(text)
        return number if math.isfinite(number) else None
    except ValueError:
        return None


def parse_datetime(value: Any) -> tuple[str | None, str | None]:
    """Return ISO timestamp and a quality issue, if one exists."""
    if value is None or value == "":
        return None, None
    if isinstance(value, datetime):
        return value.replace(tzinfo=value.tzinfo or BUSINESS_TZ).astimezone(BUSINESS_TZ).isoformat(), None
    if isinstance(value, date):
        return datetime.combine(value, datetime.min.time()).replace(tzinfo=BUSINESS_TZ).isoformat(), None

    text = str(value).strip()
    if not text:
        return None, None
    for fmt in (
        "%Y-%m-%d %H:%M:%S",
        "%Y-%m-%d %H:%M",
        "%Y-%m-%d",
        "%Y/%m/%d %H:%M:%S",
        "%Y/%m/%d %H:%M",
        "%Y/%m/%d",
    ):
        try:
            parsed = datetime.strptime(text, fmt).replace(tzinfo=BUSINESS_TZ)
            return parsed.isoformat(), None
        except ValueError:
            pass
    return None, f"非法时间值：{text}"


def parse_day(value: Any) -> tuple[str | None, str | None]:
    timestamp, issue = parse_datetime(value)
    if timestamp:
        return timestamp[:10], issue
    return None, issue


def business_dates_for_record(record: dict[str, Any]) -> list[str]:
    """Keep the +1-day pickup rule and additionally include production-day
    records through 12:00 (inclusive) in that same business date.
    """
    production_day = record.get("production_date")
    if not production_day:
        return []
    try:
        day = date.fromisoformat(production_day)
    except (TypeError, ValueError):
        return []

    result = { (day + timedelta(days=1)).isoformat() }
    production_at = record.get("production_at")
    if production_at:
        try:
            parsed = datetime.fromisoformat(production_at)
        except (TypeError, ValueError):
            parsed = None
        if parsed and parsed.timetz().replace(tzinfo=None) <= time(12, 0):
            result.add(day.isoformat())
    return sorted(result)


def channel_for(shop_name: str | None) -> str:
    shop = shop_name or ""
    for keyword, channel in CHANNEL_RULES:
        if keyword in shop:
            return channel
    return "其他平台"


def safe_value(row: tuple[Any, ...], index: dict[str, int], key: str) -> Any:
    position = index.get(key)
    return row[position] if position is not None and position < len(row) else None


def read_source(source: Path) -> tuple[list[str], list[tuple[Any, ...]], str]:
    if source.suffix.lower() == ".csv":
        decoded = None
        for encoding in ("utf-8-sig", "gb18030"):
            try:
                decoded = source.read_text(encoding=encoding)
                break
            except UnicodeDecodeError:
                continue
        if decoded is None:
            raise ValueError("CSV 编码不是 UTF-8 或 GB18030")
        reader = csv.reader(decoded.splitlines())
        rows = list(reader)
        if not rows:
            raise ValueError("空 CSV 文件")
        return [str(value).strip() for value in rows[0]], [tuple(row) for row in rows[1:]], "CSV"

    if source.suffix.lower() != ".xlsx":
        raise ValueError("仅支持真正的 .xlsx 或 .csv TMS 导出文件")
    with source.open("rb") as stream:
        if stream.read(2) != b"PK":
            raise ValueError("文件内容不是 XLSX；可能下载到了登录页")
    workbook = openpyxl.load_workbook(source, data_only=True, read_only=True)
    worksheet = workbook["SOO明细"] if "SOO明细" in workbook.sheetnames else workbook.active
    rows = worksheet.iter_rows(values_only=True)
    first = next(rows, None)
    if not first:
        raise ValueError("工作表为空")
    return [str(value).strip() if value is not None else "" for value in first], [tuple(row) for row in rows], worksheet.title


def extract(source: Path) -> dict[str, Any]:
    headers, body_rows, sheet = read_source(source)
    index: dict[str, int] = {}
    for position, name in enumerate(headers[:SOURCE_COLUMNS]):
        if name and name not in index:
            index[name] = position
    missing = REQUIRED_COLUMNS - index.keys()
    if missing:
        raise ValueError(f"缺少 TMS 必填列：{', '.join(sorted(missing))}")
    total_body_rows = len(body_rows)

    records: list[dict[str, Any]] = []
    raw_rows: list[dict[str, Any]] = []
    quality_counts = {
        "blank_soo_rows_excluded": 0,
        "invalid_datetime_values": 0,
        "invalid_link_datetime_values": 0,
        "invalid_qty_values": 0,
        "negative_qty_values": 0,
        "state_conflicts": 0,
    }
    duplicate_keys: set[str] = set()
    duplicate_rows = 0

    for row_number, row in enumerate(body_rows, start=2):
        soo_no = clean_text(safe_value(row, index, "SOO订单号"))
        if not soo_no:
            quality_counts["blank_soo_rows_excluded"] += 1
            continue

        domain_code = clean_text(safe_value(row, index, "域")) or "unknown"
        key = f"{domain_code}:{soo_no}"
        if key in duplicate_keys:
            duplicate_rows += 1
        duplicate_keys.add(key)

        flags: list[str] = []
        production_at, production_issue = parse_datetime(safe_value(row, index, "生产日期"))
        production_date = production_at[:10] if production_at else None
        load_completed_at, load_issue = parse_datetime(safe_value(row, index, "装载完成时间"))
        pickup_at, pickup_issue = parse_datetime(safe_value(row, index, "实际提货日期"))
        latest_pickup_at, latest_issue = parse_datetime(safe_value(row, index, "最晚提货时间"))
        source_updated_at, updated_issue = parse_datetime(safe_value(row, index, "更新日期"))
        tms_created_at, created_issue = parse_datetime(safe_value(row, index, "创建时间"))
        park_entered_at, entered_issue = parse_datetime(safe_value(row, index, "围栏进入时间"))
        park_exited_at, exited_issue = parse_datetime(safe_value(row, index, "围栏离开时间"))
        timeline_flags = [issue for issue in (created_issue, entered_issue, exited_issue) if issue]
        quality_counts["invalid_link_datetime_values"] += len(timeline_flags)

        for issue in (production_issue, load_issue, pickup_issue, latest_issue, updated_issue):
            if issue:
                flags.append(issue)
                quality_counts["invalid_datetime_values"] += 1

        quantity_raw = safe_value(row, index, "soo下发件数")
        quantity = parse_number(quantity_raw)
        if quantity is None:
            flags.append("件数不是有效数字")
            quality_counts["invalid_qty_values"] += 1
        elif quantity < 0:
            flags.append("件数为负数")
            quality_counts["negative_qty_values"] += 1

        source_status = clean_text(safe_value(row, index, "状态"))
        tms_status = source_status or "未知"
        is_picked_up = is_picked_up_status(source_status)
        if bool(pickup_at) != is_picked_up:
            flags.append("状态与实际提货时间冲突")
            quality_counts["state_conflicts"] += 1

        shop_name = clean_text(safe_value(row, index, "店铺名称"))
        raw_rows.append({"row_no": row_number, "domain_code": domain_code, "soo_no": soo_no,
                         "payload": {name: str(row[position]) if row[position] is not None else None
                                     for name, position in index.items() if position < len(row)}})
        record = {
            "row_no": row_number,
            "domain_code": domain_code,
            "soo_no": soo_no,
            "shipping_order_id": clean_text(safe_value(row, index, "发货订单Id")),
            "waybill_id": clean_text(safe_value(row, index, "运单ID")),
            "warehouse_code": clean_text(safe_value(row, index, "起始地ID")),
            "warehouse_name": clean_text(safe_value(row, index, "起始地名称")) or "未映射仓库",
            "carrier_code": clean_text(safe_value(row, index, "实际承运商ID")),
            "carrier_name": clean_text(safe_value(row, index, "实际承运商名称")) or "未分配承运商",
            "shop_code": clean_text(safe_value(row, index, "店铺ID")),
            "shop_name": shop_name,
            "channel": channel_for(shop_name),
            "production_date": production_date,
            "production_at": production_at,
            "tms_created_at": tms_created_at,
            "load_completed_at": load_completed_at,
            "park_entered_at": park_entered_at,
            "park_exited_at": park_exited_at,
            "actual_pickup_at": pickup_at,
            "latest_pickup_at": latest_pickup_at,
            "tms_status": tms_status,
            "is_picked_up": is_picked_up,
            "is_loaded": bool(load_completed_at),
            "soo_qty": quantity,
            "gross_cbm": parse_number(safe_value(row, index, "毛体积")),
            "gross_weight": parse_number(safe_value(row, index, "毛重")),
            "carton_count": parse_number(safe_value(row, index, "发货订单箱数")),
            "origin_province": clean_text(safe_value(row, index, "起始地省份")),
            "origin_city": clean_text(safe_value(row, index, "起始地城市")),
            "dest_province": clean_text(safe_value(row, index, "目的地省份")),
            "dest_city": clean_text(safe_value(row, index, "目的地城市")),
            "source_updated_at": source_updated_at,
            "quality_flags": flags,
            "timeline_quality_flags": timeline_flags,
        }
        records.append(record)

    business_dates = sorted({business_date for record in records
                             for business_date in business_dates_for_record(record)})
    source_updates = sorted({record["source_updated_at"] for record in records if record["source_updated_at"]})
    output = {
        "schema_version": "mvp-1.0",
        "source": {
            "file_name": source.name,
            "file_sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
            "sheet": sheet,
            "body_range": f"A1:{'BH' if len(headers) >= 60 else 'BC'}{total_body_rows + 1}",
            "last_source_updated_at": source_updates[-1] if source_updates else None,
        },
        "quality": {
            "total_body_rows": total_body_rows,
            "valid_soo_count": len(records),
            "duplicate_key_rows": duplicate_rows,
            **quality_counts,
        },
        "business_dates": business_dates,
        "records": records,
        "raw_rows": raw_rows,
    }
    return output


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", default="/Users/harumi/Downloads/提货汇总套表.xlsx")
    parser.add_argument("--output", default="control-tower/dist/data.json")
    args = parser.parse_args()
    source = Path(args.input)
    target = Path(args.output)
    target.parent.mkdir(parents=True, exist_ok=True)
    output = extract(source)
    output.pop("raw_rows")  # Never serve address, contact name or phone in dashboard JSON.
    target.write_text(json.dumps(output, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"output": str(target), "valid_soo_count": len(output["records"]), "quality": output["quality"]}, ensure_ascii=False))


if __name__ == "__main__":
    main()
