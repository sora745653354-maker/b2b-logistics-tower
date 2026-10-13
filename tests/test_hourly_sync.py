from __future__ import annotations

import os
from datetime import datetime, timedelta, timezone
from pathlib import Path

import pytest
from fastapi.testclient import TestClient

from backend import sqlite_app as tower
from ingestion import hourly_tms_sync as sync
from test_sqlite_ingestion import isolated, workbook  # noqa: F401


def fresh_report(path: Path, run: dict, quantity: int = 12) -> Path:
    source = workbook(path, f"{run['report_start'][:10]} 15:00:00", quantity)
    # Align the fixture's production date with the moving business window.
    import openpyxl
    book = openpyxl.load_workbook(source)
    book.active.cell(2, 3).value = run["report_start"][:10]
    book.active.cell(2, 4).value = f"{run['report_start'][:10]} 10:00:00"
    book.save(source)
    return source


def test_unchanged_fresh_export_updates_check_not_import(isolated: Path) -> None:
    run = sync.start_run()
    source = fresh_report(isolated / "report.xlsx", run)
    first = sync.complete_run(run["run_id"], source, 1)
    imported_at = tower.public_snapshot()["source"]["last_success_import_at"]
    run = sync.start_run()
    os.utime(source, None)
    second = sync.complete_run(run["run_id"], source, 1)
    snapshot = tower.public_snapshot()
    assert first["batch_id"] == second["batch_id"] and second["duplicate"]
    assert snapshot["source"]["last_success_import_at"] == imported_at
    assert snapshot["source"]["last_tms_check_at"] >= imported_at
    assert snapshot["source"]["last_success_sync_at"] == snapshot["collection"]["last_success_at"]


def test_old_file_cannot_claim_freshness(isolated: Path) -> None:
    run = sync.start_run()
    source = fresh_report(isolated / "old.xlsx", run)
    old_stamp = (datetime.now(timezone.utc) - timedelta(hours=2)).timestamp()
    os.utime(source, (old_stamp, old_stamp))
    with pytest.raises(ValueError, match="旧文件"):
        sync.complete_run(run["run_id"], source)
    with tower.connect() as conn:
        status = tower.collection_status(conn)
    assert status["latest_run"]["status"] == "failed"
    assert status["last_success_at"] is None
    assert tower.public_snapshot()["records"] == []


def test_count_mismatch_keeps_previous_snapshot(isolated: Path) -> None:
    first = sync.start_run()
    sync.complete_run(first["run_id"], fresh_report(isolated / "first.xlsx", first))
    before = tower.public_snapshot()
    second = sync.start_run()
    with pytest.raises(ValueError, match="不符"):
        sync.complete_run(second["run_id"], fresh_report(isolated / "bad.xlsx", second, 99), 2)
    after = tower.public_snapshot()
    assert after["records"] == before["records"]
    assert after["source"]["last_success_sync_at"] == before["source"]["last_success_sync_at"]


def test_overlapping_exports_are_rejected(isolated: Path) -> None:
    sync.start_run()
    with pytest.raises(ValueError, match="重叠"):
        sync.start_run()


def test_out_of_window_file_is_rejected(isolated: Path) -> None:
    run = sync.start_run()
    source = workbook(isolated / "wrong.xlsx", "2026-10-08 15:00:00", 3)
    with tower.connect() as conn:
        conn.execute("UPDATE tms_sync_run SET report_start='2099-01-01 00:00:00',"
                     "report_end='2099-01-03 00:00:00' WHERE run_id=?", (run["run_id"],))
    with pytest.raises(ValueError, match="范围以外"):
        sync.complete_run(run["run_id"], source)
    assert tower.public_snapshot()["records"] == []


def test_health_and_failure_are_visible(isolated: Path) -> None:
    run = sync.start_run()
    sync.fail_run(run["run_id"], "auth_required", "需要完成 TMS 登录")
    with TestClient(tower.app) as client:
        status = client.get("/api/v1/sync/status").json()
        data = client.get("/data.json")
    assert status["latest_run"]["error_code"] == "auth_required"
    assert data.headers["Cache-Control"] == "no-store"
    assert data.json()["collection"]["last_success_at"] is None


def test_fallback_watcher_cannot_publish_unverified_hourly_file(isolated: Path, monkeypatch) -> None:
    from ingestion import watch_tms_exports as watcher
    monkeypatch.setattr(watcher, "DOWNLOAD_DIR", isolated)
    run = sync.start_run()
    source = fresh_report(isolated / "订单明细报表（soo）.xlsx", run)
    # File is stable enough for the watcher but belongs to the current run.
    stamp = datetime.fromisoformat(run["started_at"]).timestamp() - 4
    os.utime(source, (stamp, stamp))
    assert watcher.scan_once() == []
    with pytest.raises(ValueError, match="查询总数"):
        sync.complete_run(run["run_id"], source, expected_rows=2)
    assert watcher.scan_once() == []
    assert tower.public_snapshot()["records"] == []
