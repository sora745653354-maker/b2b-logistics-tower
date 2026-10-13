from __future__ import annotations

import json
from pathlib import Path

import openpyxl
import pytest
from fastapi.testclient import TestClient

from backend import sqlite_app as tower
from ingestion.build_origin_routes import carrier_map, load_routes
from ingestion.build_warehouse_regions import load_warehouse_regions, write_catalog


def test_route_workbook_is_converted_to_minimal_catalog(tmp_path: Path) -> None:
    source = tmp_path / "routes.xlsx"
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.append(["始发仓库", "物流公司", "收货城市", "收货省份", "始发区域"])
    sheet.append(["测试仓", "测试物流", "苏州", "江苏", "华东"])
    book.save(source)

    assert load_routes(source) == [{"warehouse": "测试仓", "carrier": "测试物流",
                                    "city": "苏州", "province": "江苏", "region": "华东"}]


def test_route_workbook_rejects_unexpected_headers(tmp_path: Path) -> None:
    source = tmp_path / "wrong.xlsx"
    book = openpyxl.Workbook()
    book.active.append(["仓库", "承运商"])
    book.save(source)

    with pytest.raises(ValueError, match="表头不符合预期"):
        load_routes(source)


def test_carrier_map_groups_by_origin_without_duplicates() -> None:
    assert carrier_map([
        {"warehouse": "华东光合仓", "carrier": "物流甲"},
        {"warehouse": "华东光合仓", "carrier": "物流乙"},
        {"warehouse": "华东光合仓", "carrier": "物流甲"},
    ]) == {"华东光合仓": ["物流乙", "物流甲"]}


def test_route_catalog_endpoint_serves_json(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    dist = tmp_path / "dist"
    dist.mkdir()
    (dist / "warehouse-carriers.json").write_text(
        json.dumps({"warehouse_carriers": {"测试仓": ["测试物流"]},
                    "warehouse_regions": {"测试仓": "华东仓"}}), encoding="utf-8")
    monkeypatch.setattr(tower, "ROOT", tmp_path)

    with TestClient(tower.app) as client:
        response = client.get("/warehouse-carriers.json")

    assert response.status_code == 200
    assert response.json()["warehouse_carriers"] == {"测试仓": ["测试物流"]}
    assert response.json()["warehouse_regions"] == {"测试仓": "华东仓"}
    assert response.headers["cache-control"] == "no-store"


def test_outbound_catalog_endpoint_serves_json(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    dist = tmp_path / "dist"
    dist.mkdir()
    expected = {"source_file": "outbound.xlsx", "record_count": 1,
                "records": [{"soo_no": "SOO-1", "actual_shipment_at": "2026-10-10 12:00:00"}]}
    (dist / "outbound-data.json").write_text(json.dumps(expected), encoding="utf-8")
    monkeypatch.setattr(tower, "ROOT", tmp_path)

    with TestClient(tower.app) as client:
        response = client.get("/outbound-data.json")

    assert response.status_code == 200
    assert response.json() == expected
    assert response.headers["cache-control"] == "no-store"


def test_route_catalog_includes_all_warehouse_regions() -> None:
    catalog_path = tower.ROOT / "dist" / "warehouse-carriers.json"
    catalog = json.loads(catalog_path.read_text(encoding="utf-8"))
    regions = catalog["warehouse_regions"]

    assert len(regions) == 21
    assert regions["华东光合仓"] == "华东光合仓"
    assert regions["湖州优全VMI仓"] == "优全工厂仓"
    assert regions["滕州-滕州茁彼VMI仓"] == "茁彼工厂仓"


def test_warehouse_region_workbook_supports_leading_empty_column(tmp_path: Path) -> None:
    source = tmp_path / "warehouse-regions.xlsx"
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.append([None, "仓库名称", "仓区域"])
    sheet.append([None, "华东光合仓", "华东光合仓"])
    book.save(source)

    assert load_warehouse_regions(source) == {"华东光合仓": "华东光合仓"}


def test_warehouse_region_catalog_update_preserves_carriers(tmp_path: Path) -> None:
    catalog = tmp_path / "warehouse-carriers.json"
    catalog.write_text(json.dumps({"warehouse_carriers": {"测试仓": ["测试物流"]}}), encoding="utf-8")

    write_catalog({"测试仓": "华东仓"}, catalog)

    saved = json.loads(catalog.read_text(encoding="utf-8"))
    assert saved["warehouse_carriers"] == {"测试仓": ["测试物流"]}
    assert saved["warehouse_regions"] == {"测试仓": "华东仓"}
